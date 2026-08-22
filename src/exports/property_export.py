"""Inventory export: properties → a spreadsheet the admin can download.

Columns, labels and formats are Spanish and mirror what the back-office shows in
the properties table. Two serializations of the same rows:

* ``to_xlsx`` — a real Excel file (numbers stay numbers, so totals/filters work).
* ``to_csv``  — UTF-8 **with BOM** and ``;`` as the separator, the combination
  Excel in a Spanish locale opens straight into columns (decimals use ``,``).
"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any, Callable, NamedTuple

from src.db.models.property_models import PropertyORM

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"

_STATUS = {
    "draft": "Borrador",
    "pending": "Pendiente",
    "published": "Publicado",
    "paused": "Pausado",
    "rejected": "Rechazado",
    "sold": "Vendido",
    "rented": "Arrendado",
    "deleted": "Eliminado",
}
_OPERATION = {"sale": "Venta", "rent": "Arriendo", "temporary": "Temporal"}
_KIND = {
    "house": "Casa",
    "apartment": "Apartamento",
    "studio": "Apartaestudio",
    "lot": "Lote",
    "office": "Oficina",
    "commercial": "Local comercial",
    "warehouse": "Bodega",
    "farm": "Finca",
}
_CONDITION = {
    "new": "Nuevo",
    "used": "Usado",
    "remodeled": "Remodelado",
    "under_construction": "En construcción",
}
_VIEW = {"internal": "Interna", "external": "Externa"}
_SECURITY = {"none": "Sin vigilancia", "private": "Privada (portería)", "automated": "Automatizada"}


class ExportContext(NamedTuple):
    """Data shared by every row: resolved once instead of per property."""

    city_names: dict[str, str]  # location_id → nearest city name
    site_url: str               # public base URL, for the listing link column


# ---------------------------------------------------------------------------
# Cell helpers
# ---------------------------------------------------------------------------

def _label(mapping: dict[str, str], value: str | None) -> str:
    """Spanish label for a code, falling back to the raw code (never blank-out data)."""
    if not value:
        return ""
    return mapping.get(value, value)


def _money(minor_units: int | None) -> int | float | str:
    """Minor units (centavos) → major units, as a number Excel can sum."""
    if minor_units is None:
        return ""
    major = minor_units / 100
    return int(major) if major == int(major) else major


def _yes_no(value: bool | None) -> str:
    return "Sí" if value else "No"


def _day(value: datetime | date | None) -> str:
    if value is None:
        return ""
    return value.strftime("%Y-%m-%d")


def _city(prop: PropertyORM, ctx: ExportContext) -> str:
    return ctx.city_names.get(prop.location_id or "", "")


def _zone(prop: PropertyORM) -> str:
    """The location node the property hangs off (barrio/localidad, or the city)."""
    return prop.location.name if prop.location else ""


def _url(prop: PropertyORM, ctx: ExportContext) -> str:
    return f"{ctx.site_url}/propiedades/{prop.nid}" if ctx.site_url else ""


# ---------------------------------------------------------------------------
# Columns — the sheet layout, in reading order
# ---------------------------------------------------------------------------

_Getter = Callable[[PropertyORM, ExportContext], Any]

COLUMNS: tuple[tuple[str, _Getter], ...] = (
    ("ID", lambda p, c: p.nid),
    ("Título", lambda p, c: p.title or ""),
    ("Estado", lambda p, c: _label(_STATUS, p.status)),
    ("Operación", lambda p, c: _label(_OPERATION, p.operation_type)),
    ("Tipo", lambda p, c: _label(_KIND, p.property_kind)),
    ("Estado del inmueble", lambda p, c: _label(_CONDITION, p.condition)),
    ("Precio", lambda p, c: _money(p.price_amount)),
    ("Moneda", lambda p, c: p.currency or ""),
    ("Administración", lambda p, c: _money(p.admin_fee_amount)),
    ("Área total (m²)", lambda p, c: p.total_area_m2 if p.total_area_m2 is not None else ""),
    ("Área construida (m²)", lambda p, c: p.built_area_m2 if p.built_area_m2 is not None else ""),
    ("Habitaciones", lambda p, c: p.bedrooms if p.bedrooms is not None else ""),
    ("Baños", lambda p, c: p.bathrooms if p.bathrooms is not None else ""),
    ("Parqueaderos", lambda p, c: p.parking_spots if p.parking_spots is not None else ""),
    ("Estrato", lambda p, c: p.stratum if p.stratum is not None else ""),
    ("Piso", lambda p, c: p.floor_number if p.floor_number is not None else ""),
    ("Pisos del edificio", lambda p, c: p.total_floors if p.total_floors is not None else ""),
    ("Antigüedad (años)", lambda p, c: p.age_years if p.age_years is not None else ""),
    ("Vista", lambda p, c: _label(_VIEW, p.view_type)),
    ("Vigilancia", lambda p, c: _label(_SECURITY, p.security_type)),
    ("Depósito", lambda p, c: _yes_no(p.has_storage)),
    ("Ascensor", lambda p, c: _yes_no(p.has_elevator)),
    ("Estudio", lambda p, c: _yes_no(p.has_study)),
    ("Balcón", lambda p, c: _yes_no(p.has_balcony)),
    ("Ciudad", _city),
    ("Zona / Barrio", lambda p, c: _zone(p)),
    ("Dirección", lambda p, c: p.address_street or ""),
    ("Detalle dirección", lambda p, c: p.address_detail or ""),
    ("Propietario", lambda p, c: (p.owner.full_name if p.owner else "") or ""),
    ("Email propietario", lambda p, c: (p.owner.email if p.owner else "") or ""),
    ("Inmobiliaria", lambda p, c: (p.agency.name if p.agency else "") or ""),
    ("Teléfono contacto", lambda p, c: p.contact_phone or ""),
    ("WhatsApp", lambda p, c: p.contact_whatsapp or ""),
    ("Email contacto", lambda p, c: p.contact_email or ""),
    ("Fotos", lambda p, c: len(p.images or [])),
    ("Vistas", lambda p, c: p.views_count or 0),
    ("Leads", lambda p, c: p.leads_count or 0),
    ("Favoritos", lambda p, c: p.favorites_count or 0),
    ("En home", lambda p, c: _yes_no(p.show_on_home)),
    ("Creado", lambda p, c: _day(p.created_at)),
    ("Publicado", lambda p, c: _day(p.published_at)),
    ("Vence", lambda p, c: _day(p.expires_at)),
    ("Motivo de rechazo", lambda p, c: p.rejection_reason or ""),
    ("URL", _url),
)

HEADERS: tuple[str, ...] = tuple(header for header, _ in COLUMNS)


def build_rows(
    properties: list[PropertyORM],
    *,
    city_names: dict[str, str] | None = None,
    site_url: str = "",
) -> list[list[Any]]:
    """One list of cell values per property, in ``COLUMNS`` order."""
    ctx = ExportContext(city_names=city_names or {}, site_url=site_url.rstrip("/"))
    return [[getter(prop, ctx) for _, getter in COLUMNS] for prop in properties]


# ---------------------------------------------------------------------------
# Serialization
# ---------------------------------------------------------------------------

def _csv_cell(value: Any) -> str:
    """Decimals use ``,`` — the separator a Spanish-locale Excel reads as decimal."""
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.2f}".rstrip("0").rstrip(".").replace(".", ",")
    return str(value)


def to_csv(rows: list[list[Any]], headers: tuple[str, ...] = HEADERS) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_csv_cell(cell) for cell in row])
    # BOM: without it Excel reads UTF-8 accents as mojibake (Baños → BaÃ±os).
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def to_xlsx(rows: list[list[Any]], headers: tuple[str, ...] = HEADERS) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventario"

    sheet.append(list(headers))
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="1D4ED8")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 22
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    # Width from the widest cell in each column (capped so a long description
    # cannot push the rest of the sheet off screen).
    for index, header in enumerate(headers, start=1):
        widest = max([len(header)] + [len(str(row[index - 1])) for row in rows])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(widest + 2, 10), 45)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
